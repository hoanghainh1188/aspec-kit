# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1,<4"]
# ///
"""Generate the sample 詳細設計 workbook used by the converter tests.

    uv run --script packages/aspec-kit/test/convert/make_fixture.py

Writes `fixtures/detail-design-sample.xlsx` (committed). It mimics a typical Japanese client workbook:
Japanese sheet names, a cover sheet, a 方眼紙-style header with merged cells, a field/validation table
with vertically merged group cells, an error-message sheet, plus a sheet name with characters that are
unsafe in file names.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent / "fixtures" / "detail-design-sample.xlsx"
FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)


def cover(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "表紙"
    ws["B2"] = "詳細設計書"
    ws.merge_cells("B2:F2")
    ws["B4"] = "システム名"
    ws["D4"] = "予約管理システム"
    ws.merge_cells("B4:C4")
    ws.merge_cells("D4:F4")
    ws["B5"] = "作成日"
    ws["D5"] = dt.datetime(2026, 7, 10)
    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:F5")


def field_spec(wb: Workbook) -> None:
    ws = wb.create_sheet("画面項目定義")
    ws["A1"] = "画面項目定義（予約登録画面）"
    ws.merge_cells("A1:G1")
    # 方眼紙 header: 項目名 spans B:C, 桁数 is a plain cell
    header = {"A3": "No", "B3": "項目名", "D3": "型", "E3": "桁数", "F3": "必須", "G3": "バリデーション"}
    for ref, value in header.items():
        ws[ref] = value
    ws.merge_cells("B3:C3")
    rows = [
        (1, "予約者名", "文字列", 50, "○", "全角のみ。空白不可"),
        (2, "メールアドレス", "文字列", 254, "", "RFC 5322 形式 | 重複不可"),
        (3, "予約日", "日付", 10, "○", "本日以降\n営業日のみ"),
        (4, "人数", "数値", 2, "", "1〜20"),
    ]
    for i, (no, name, typ, digits, req, rule) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=no)
        ws.cell(row=i, column=2, value=name)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        ws.cell(row=i, column=4, value=typ)
        ws.cell(row=i, column=5, value=digits)
        ws.cell(row=i, column=6, value=req or None)
        ws.cell(row=i, column=7, value=rule)
    # Vertical merge: rows 4-5 share one 必須 cell (row 5 itself is empty → value must be repeated)
    ws.merge_cells("F4:F5")


def messages(wb: Workbook) -> None:
    ws = wb.create_sheet("エラーメッセージ")
    ws.append(["コード", "メッセージ", "表示位置"])
    ws.append(["E001", "予約者名を入力してください。", "項目下"])
    ws.append(["E002", "メールアドレスの形式が正しくありません。", "項目下"])
    ws.append(["E003", "予約日は本日以降を指定してください。", "画面上部"])


def unsafe_name(wb: Workbook) -> None:
    ws = wb.create_sheet("改訂履歴 <v1|v2>")  # <, > and | are legal in sheet names but not in file names
    ws.append(["版", "日付", "内容"])
    ws.append(["1.0", dt.datetime(2026, 7, 1), "初版"])
    ws.append(["1.1", dt.datetime(2026, 7, 10), "人数の上限を 20 に変更"])


def main() -> None:
    wb = Workbook()
    cover(wb)
    field_spec(wb)
    messages(wb)
    unsafe_name(wb)
    wb.properties.creator = "aspec-kit"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
