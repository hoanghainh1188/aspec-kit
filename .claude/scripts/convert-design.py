#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "openpyxl>=3.1,<4",
#   "pdfplumber>=0.11,<1",
# ]
# ///
"""Convert a 詳細設計 (detail design) workbook or PDF into Markdown that Claude Code can read.

    uv run --script .claude/scripts/convert-design.py <file.xlsx|file.xlsm|file.pdf> [--out DIR] [--force]
    npx @hoanghainh/aspec-kit convert <file> [--out DIR] [--force]      # same thing, via the CLI

- The original file is never modified (it stays the source of truth). Markdown is written next to it.
- Excel: `<stem>.md` (index) + one `<stem>__<NN>_<sheet>.md` per sheet. Merged cells are expanded
  (the value is repeated in every cell of the range) and "Excel 方眼紙" column spans are collapsed.
- PDF: `<stem>.md` with text + tables per page.
- Every run appends an entry to `<out>/CHANGELOG.md`.
- A Markdown file is only overwritten if this tool generated it from the same source (or --force).

Default --out: the `docs/02-detail-design/<slug>/` folder that contains the original; if the original
lives elsewhere, `./docs/02-detail-design/<slug-of-file-name>/`.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

DETAIL_DESIGN_DIR = ("docs", "02-detail-design")
EXCEL_EXTS = {".xlsx", ".xlsm"}
PDF_EXTS = {".pdf"}
MARKER_PREFIX = "<!-- aspec-kit convert: source="
MAX_LISTED_MERGES = 40

# Characters that are unsafe in file names on Windows/macOS/Linux, plus control characters.
UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]')


class ConvertError(Exception):
    """User-facing error (printed without a traceback)."""


# ─────────────────────────────────────────────────────────────────────────────
# Naming helpers
# ─────────────────────────────────────────────────────────────────────────────


def sanitize_filename(name: str, fallback: str = "sheet") -> str:
    """Keep the name (incl. Japanese) as-is; replace only characters unsafe in file names."""
    cleaned = UNSAFE_FILENAME_RE.sub("_", unicodedata.normalize("NFC", name))
    cleaned = cleaned.strip().rstrip(". ")
    return cleaned or fallback


def slugify(name: str) -> str:
    """Folder slug: sanitized, whitespace → '-', ASCII lower-cased (Japanese kept)."""
    slug = re.sub(r"\s+", "-", sanitize_filename(name, fallback="design")).lower()
    return re.sub(r"-{2,}", "-", slug).strip("-") or "design"


def marker_line(source_name: str) -> str:
    return f"{MARKER_PREFIX}{source_name} — generated file, do not edit; re-run the converter -->"


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ─────────────────────────────────────────────────────────────────────────────
# Markdown helpers
# ─────────────────────────────────────────────────────────────────────────────


def cell_text(value: object) -> str:
    """Render a cell value as single-line Markdown-safe text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    text = text.replace("|", "\\|")
    return "<br>".join(line.strip() for line in text.split("\n"))


def gfm_table(rows: list[list[str]]) -> str:
    """First row is the header. Rows are padded to the same width."""
    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]
    header, *body = padded
    lines = [
        "| " + " | ".join(header) + " |",
        "|" + "|".join(["---"] * width) + "|",
    ]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Cell:
    text: str
    merge_id: int | None  # index of the merged range this cell belongs to (None = not merged)


def sheet_grid(ws) -> tuple[list[list[Cell]], list[str]]:
    """Return the sheet as a grid with merged ranges expanded, plus the merged range refs."""
    if ws.max_row is None or ws.max_column is None:
        return [], []
    min_row, min_col = ws.min_row, ws.min_column
    rows = [
        [cell_text(c.value) for c in row]
        for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, min_col=min_col, max_col=ws.max_column)
    ]
    merge_ids: list[list[int | None]] = [[None] * len(r) for r in rows]
    refs: list[str] = []
    for idx, rng in enumerate(sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))):
        refs.append(rng.coord)
        top_left = cell_text(ws.cell(row=rng.min_row, column=rng.min_col).value)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                ri, ci = r - min_row, c - min_col
                if 0 <= ri < len(rows) and 0 <= ci < len(rows[ri]):
                    rows[ri][ci] = top_left
                    merge_ids[ri][ci] = idx
    grid = [[Cell(t, m) for t, m in zip(r, ids)] for r, ids in zip(rows, merge_ids)]
    return grid, refs


def split_blocks(grid: list[list[Cell]]) -> list[list[list[Cell]]]:
    """Split the grid into blocks of consecutive non-empty rows."""
    blocks: list[list[list[Cell]]] = []
    current: list[list[Cell]] = []
    for row in grid:
        if any(c.text for c in row):
            current.append(row)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def drop_empty_columns(block: list[list[Cell]]) -> list[list[Cell]]:
    width = max(len(r) for r in block)
    keep = [c for c in range(width) if any(c < len(r) and r[c].text for r in block)]
    return [[r[c] if c < len(r) else Cell("", None) for c in keep] for r in block]


def collapse_merged_spans(block: list[list[Cell]]) -> list[list[Cell]]:
    """Merge adjacent columns that only differ by a horizontal merge (Excel 方眼紙 layouts).

    Column c+1 is folded into column c when, in every row, both cells belong to the same merged
    range or are both empty. The merged value is already repeated, so nothing is lost.
    """
    columns = [list(col) for col in zip(*block)]
    if not columns:
        return block
    kept = [columns[0]]
    for col in columns[1:]:
        prev = kept[-1]
        same_span = all(
            (a.merge_id is not None and a.merge_id == b.merge_id) or (not a.text and not b.text)
            for a, b in zip(prev, col)
        )
        if not same_span:
            kept.append(col)
    return [list(row) for row in zip(*kept)]


def render_block(block: list[list[Cell]]) -> str:
    block = collapse_merged_spans(drop_empty_columns(block))
    texts = [[c.text for c in row] for row in block]
    parts: list[str] = []
    # Leading rows holding a single value (titles, notes) are rendered as text, not table rows.
    width = len(texts[0]) if texts else 0
    while len(texts) > 1 and width > 1 and len({t for t in texts[0] if t}) == 1:
        parts.append(f"**{next(t for t in texts[0] if t)}**")
        texts = texts[1:]
    if len(texts) >= 2 and width >= 2:
        parts.append(gfm_table(texts))
    else:
        parts.extend(" | ".join(t for t in row if t) for row in texts)
    return "\n\n".join(parts)


def convert_sheet(ws, source_name: str, sha: str, position: str) -> str:
    grid, merged_refs = sheet_grid(ws)
    out = [
        marker_line(source_name),
        f"# {ws.title}",
        "",
        f"> Source: `{source_name}` › sheet `{ws.title}` ({position}) · sha256 `{sha[:12]}`"
        + (" · hidden sheet" if ws.sheet_state != "visible" else ""),
        "",
    ]
    blocks = [b for b in (render_block(block) for block in split_blocks(grid)) if b]
    out.extend(b + "\n" for b in blocks)
    if not blocks:
        out.append("_(empty sheet)_\n")
    if merged_refs:
        listed = ", ".join(f"`{r}`" for r in merged_refs[:MAX_LISTED_MERGES])
        more = len(merged_refs) - MAX_LISTED_MERGES
        out.append(
            f"> Merged cells (value repeated across each range): {listed}"
            + (f" … and {more} more" if more > 0 else "")
            + "\n"
        )
    return "\n".join(out).rstrip("\n") + "\n"


def convert_excel(src: Path, sha: str) -> dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - only without uv
        raise ConvertError("openpyxl is missing — run this script with `uv run --script`.") from exc
    try:
        # data_only=True → cached formula results (what the client sees in Excel)
        wb = load_workbook(src, data_only=True)
    except Exception as exc:  # openpyxl raises many types (zip, xml, encryption)
        raise ConvertError(f"cannot read workbook {src.name}: {exc}") from exc

    stem = sanitize_filename(src.stem, fallback="design")
    sheets = wb.worksheets
    files: dict[str, str] = {}
    index_rows = [["#", "Sheet", "File"]]
    used: set[str] = set()
    for i, ws in enumerate(sheets, start=1):
        base = f"{stem}__{i:02d}_{sanitize_filename(ws.title)}"
        name = f"{base}.md"
        n = 2
        while name.lower() in used:  # sanitized names may collide on case-insensitive filesystems
            name = f"{base}-{n}.md"
            n += 1
        used.add(name.lower())
        files[name] = convert_sheet(ws, src.name, sha, f"{i}/{len(sheets)}")
        index_rows.append([str(i), cell_text(ws.title), f"[{name}](<{name}>)"])

    index = [
        marker_line(src.name),
        f"# {src.stem}",
        "",
        f"> Converted from `{src.name}` (sha256 `{sha[:12]}`) by `aspec-kit convert`. "
        "The original file is the source of truth — do not edit this Markdown by hand.",
        "",
        gfm_table(index_rows),
        "",
    ]
    files[f"{stem}.md"] = "\n".join(index)
    return files


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────


def _outside(bboxes):
    def keep(obj) -> bool:
        x0, top, x1, bottom = obj.get("x0"), obj.get("top"), obj.get("x1"), obj.get("bottom")
        if None in (x0, top, x1, bottom):
            return True
        cx, cy = (x0 + x1) / 2, (top + bottom) / 2
        return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in bboxes)

    return keep


def convert_pdf(src: Path, sha: str) -> dict[str, str]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - only without uv
        raise ConvertError("pdfplumber is missing — run this script with `uv run --script`.") from exc

    out = [
        marker_line(src.name),
        f"# {src.stem}",
        "",
        f"> Converted from `{src.name}` (sha256 `{sha[:12]}`) by `aspec-kit convert`. "
        "The original file is the source of truth — do not edit this Markdown by hand.",
        "",
    ]
    try:
        pdf = pdfplumber.open(src)
    except Exception as exc:
        raise ConvertError(f"cannot read PDF {src.name}: {exc}") from exc
    with pdf:
        for number, page in enumerate(pdf.pages, start=1):
            out.append(f"## Page {number}\n")
            tables = page.find_tables()
            text = (page.filter(_outside([t.bbox for t in tables])).extract_text() or "").strip()
            if text:
                out.append(text + "\n")
            for t_index, table in enumerate(tables, start=1):
                rows = [[cell_text(c) for c in row] for row in table.extract() if any(row)]
                if rows:
                    out.append(f"**Table {number}.{t_index}**\n")
                    out.append(gfm_table(rows) + "\n")
            if not text and not tables:
                out.append("_(no extractable text on this page — scanned image? OCR is not supported)_\n")
    return {f"{sanitize_filename(src.stem, fallback='design')}.md": "\n".join(out)}


# ─────────────────────────────────────────────────────────────────────────────
# Output
# ─────────────────────────────────────────────────────────────────────────────


def default_out_dir(src: Path, cwd: Path) -> Path:
    parts = src.parent.parts
    for i in range(len(parts) - 1):
        if parts[i : i + 2] == DETAIL_DESIGN_DIR and i + 2 < len(parts):
            return Path(*parts[: i + 3])
    return cwd.joinpath(*DETAIL_DESIGN_DIR, slugify(src.stem))


def check_overwrites(out_dir: Path, files: dict[str, str], source_name: str, force: bool) -> None:
    blocked = []
    for name in files:
        target = out_dir / name
        if not target.exists() or force:
            continue
        with target.open(encoding="utf-8", errors="replace") as f:
            first = f.readline().rstrip("\n")
        if first != marker_line(source_name):
            blocked.append(name)
    if blocked:
        raise ConvertError(
            "refusing to overwrite files not generated from "
            f"{source_name}: {', '.join(blocked)} (in {out_dir}). Use --force to overwrite."
        )


def append_changelog(out_dir: Path, src: Path, sha: str, written: list[str], today: str) -> Path:
    changelog = out_dir / "CHANGELOG.md"
    try:
        source_ref = os.path.relpath(src, out_dir)
    except ValueError:  # different drive on Windows
        source_ref = str(src)
    entry = [
        f"## Converted `{src.name}` → Markdown — {today}",
        "- Tool: `aspec-kit convert` (machine conversion; the original file is unchanged and "
        "remains the source of truth)",
        f"- Source: `{Path(source_ref).as_posix()}` (sha256 `{sha[:12]}`)",
        "- Output: " + ", ".join(f"`{name}`" for name in written),
        "- **Affected issues:** none — format conversion only, no content change",
    ]
    if changelog.exists():
        existing = changelog.read_text(encoding="utf-8")
        prefix = "" if existing.endswith("\n\n") else ("\n" if existing.endswith("\n") else "\n\n")
        with changelog.open("a", encoding="utf-8") as f:
            f.write(prefix + "\n".join(entry) + "\n")
    else:
        changelog.write_text(f"# CHANGELOG — {out_dir.name}\n\n" + "\n".join(entry) + "\n", encoding="utf-8")
    return changelog


def convert(src: Path, out_dir: Path | None, force: bool = False, cwd: Path | None = None) -> list[Path]:
    src = src.resolve()
    if not src.is_file():
        raise ConvertError(f"file not found: {src}")
    ext = src.suffix.lower()
    if ext == ".xls":
        raise ConvertError("legacy .xls is not supported — open it in Excel and save as .xlsx first.")
    if ext not in EXCEL_EXTS | PDF_EXTS:
        raise ConvertError(f"unsupported file type '{src.suffix}' (expected .xlsx, .xlsm or .pdf).")

    out_dir = (out_dir or default_out_dir(src, (cwd or Path.cwd()).resolve())).resolve()
    sha = sha256_of(src)
    files = convert_excel(src, sha) if ext in EXCEL_EXTS else convert_pdf(src, sha)
    if any((out_dir / name).resolve() == src for name in files):
        raise ConvertError("output would overwrite the original file — choose another --out.")
    check_overwrites(out_dir, files, src.name, force)

    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for name, content in files.items():
        (out_dir / name).write_text(content, encoding="utf-8", newline="\n")
        written.append(out_dir / name)
    append_changelog(out_dir, src, sha, list(files), dt.date.today().isoformat())

    if sha256_of(src) != sha:  # paranoia: the source of truth must never change
        raise ConvertError(f"the original file changed during conversion: {src}")
    return written


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="convert-design",
        description="Convert a 詳細設計 .xlsx/.xlsm/.pdf into Markdown next to the original.",
    )
    parser.add_argument("file", type=Path, help="workbook (.xlsx/.xlsm) or PDF to convert")
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="output folder (default: the docs/02-detail-design/<slug>/ folder holding the file)",
    )
    parser.add_argument("--force", action="store_true", help="overwrite existing Markdown files")
    args = parser.parse_args(argv)

    try:
        written = convert(args.file, args.out, force=args.force)
    except ConvertError as exc:
        print(f"convert-design: error: {exc}", file=sys.stderr)
        return 1

    out_dir = written[0].parent
    src = args.file.resolve()
    print(f"✓ Converted {src.name} → {out_dir}")
    for path in written:
        print(f"   - {path.name}")
    print("   - CHANGELOG.md (entry appended)")
    if src.parent != out_dir:
        print(
            f"ℹ The original is not in {out_dir}. Consider copying it there so the source of truth "
            "sits next to its Markdown."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
